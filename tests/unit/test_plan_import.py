"""plan_import.extract_text 单元测试：txt/md/docx 抽取与异常路径。"""

from io import BytesIO

import pytest
from docx import Document

from travel_agent.services.plan_import import ImportError_, extract_text


def test_extract_text_plain_txt_and_md() -> None:
    text = "Day1: 昆明滇池\nDay2: 石林"
    assert extract_text("trip.txt", text.encode()) == text
    assert extract_text("trip.md", text.encode()) == text


def test_extract_text_rejects_unsupported_suffix() -> None:
    with pytest.raises(ImportError_, match="暂不支持"):
        extract_text("trip.exe", b"payload")


def test_extract_text_docx_paragraphs_and_table() -> None:
    doc = Document()
    doc.add_paragraph("第一天：滇池")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "预算"
    table.rows[0].cells[1].text = "3000 元"
    buf = BytesIO()
    doc.save(buf)

    text = extract_text("trip.docx", buf.getvalue())
    assert "第一天：滇池" in text
    assert "预算 | 3000 元" in text


def test_extract_text_docx_garbled_raises() -> None:
    with pytest.raises(ImportError_, match="Word"):
        extract_text("trip.docx", b"\x00\x01not-a-docx")


def test_extract_text_pdf_garbled_raises() -> None:
    with pytest.raises(ImportError_, match="PDF"):
        extract_text("trip.pdf", b"not-a-real-pdf")
